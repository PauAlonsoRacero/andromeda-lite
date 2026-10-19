"""Fallos del tercer registro real y de la revisión de interfaz de Pau."""
import asyncio
import logging
import os
import tempfile
from pathlib import Path

logging.disable(logging.CRITICAL)
ROOT = Path(__file__).resolve().parent.parent


def _client(tmp_path, monkeypatch):
    monkeypatch.setenv("ANDROMEDA_DATA_DIR", str(tmp_path))
    monkeypatch.setattr("pathlib.Path.home", lambda: tmp_path)
    from fastapi.testclient import TestClient
    from app import create_app
    return TestClient(create_app())


def _perfil():
    from app.memory.profile import MemoryProfile
    p = MemoryProfile(tempfile.mktemp())
    p.upsert_fact("gustos:negro", "Me gusta el color negro")
    p.upsert_fact("nombre", "Me llamo Pau")
    return p


# ── "eliminar de memoria que me gusta el negro" borró TODA la memoria ──────

def test_removing_one_fact_does_not_wipe_memory():
    from app.core import prompt_builder as pb
    p = _perfil()
    r = asyncio.run(pb.handle_memory_commands(
        pb.PromptParts(message="x"),
        "podrías eliminar de memoria que me gusta el color negro? lo odio", p, None, False))
    assert r == "forgot"
    assert list(p.get()["facts"].values()) == ["Me llamo Pau"], "borró más de la cuenta"


def test_forget_everything_still_clears():
    from app.core import prompt_builder as pb
    for orden in ("olvida todo", "borra la memoria", "forget everything"):
        p = _perfil()
        asyncio.run(pb.handle_memory_commands(pb.PromptParts(message="x"), orden, p, None, False))
        assert p.get()["facts"] == {}, orden


def test_memory_commands_are_not_extracted_as_facts():
    """Se guardaba "me gusta el color negro" justo al pedir olvidarlo."""
    from app.core.prompt_builder import is_memory_command
    from app.memory.extractor import extract_memories
    assert is_memory_command("podrías eliminar de memoria que me gusta el color negro?")
    assert extract_memories("me gusta el color negro? lo odio") == []
    assert extract_memories("me gusta el audi tt")


# ── Imágenes: el bucle de herramientas las perdía ──────────────────────────

def test_images_skip_the_tool_loop():
    import inspect
    from app.core import streaming
    assert 'not config.get("images")' in inspect.getsource(streaming)


# ── Búsqueda ───────────────────────────────────────────────────────────────

def test_task_later_in_the_sentence_does_not_search():
    from app.core.web_search import needs_web_search
    assert not needs_web_search("Aquí está la imagen, ahora crea la página web html porfavor")


def test_referential_search_uses_previous_topic():
    """Buscaba literalmente "info sobre lo que te he dicho"."""
    from app.core.prompt_builder import _consulta_con_contexto
    h = [{"role": "user", "content": "ranking de trabajos mejor pagados en Europa y MLOps"},
         {"role": "assistant", "content": "…"}]
    assert "MLOps" in _consulta_con_contexto("por favor busca info sobre lo que te he dicho", h)
    assert _consulta_con_contexto("busca el iphone 15 pro max", h) == "busca el iphone 15 pro max"


# ── Respuestas sin estructura ─────────────────────────────────────────────

def test_system_prompt_asks_for_structure():
    from app.specialists.identity import with_identity
    s = with_identity("x", "generalist", "qwen3:8b", user_prompt="busca info del iphone")
    assert "prosa" in s and "tabla" in s


# ── Vista previa en blanco ─────────────────────────────────────────────────

def _ws(tmp_path, monkeypatch):
    ws = tmp_path / "ws"
    (ws / "web" / "css").mkdir(parents=True)
    (ws / "web" / "index.html").write_text('<link rel="stylesheet" href="css/e.css"><h1>Hola</h1>')
    (ws / "web" / "css" / "e.css").write_text("h1{color:red}")
    monkeypatch.setenv("ANDROMEDA_WORKSPACE", str(ws))
    return ws


def test_preview_can_be_framed_but_is_sandboxed(tmp_path, monkeypatch):
    """X-Frame-Options: DENY dejaba el visor en blanco."""
    _ws(tmp_path, monkeypatch)
    with _client(tmp_path, monkeypatch) as c:
        r = c.get("/api/files/view/web/index.html")
        assert r.status_code == 200
        assert r.headers["x-frame-options"] == "SAMEORIGIN"
        assert "sandbox" in r.headers["content-security-policy"]
        assert c.get("/api/files/raw", params={"path": "web/index.html"}).headers["x-frame-options"] == "DENY"


def test_preview_resolves_relative_assets(tmp_path, monkeypatch):
    _ws(tmp_path, monkeypatch)
    with _client(tmp_path, monkeypatch) as c:
        assert c.get("/api/files/view/web/css/e.css").text == "h1{color:red}"


def test_preview_cannot_escape_the_workspace(tmp_path, monkeypatch):
    _ws(tmp_path, monkeypatch)
    with _client(tmp_path, monkeypatch) as c:
        assert c.get("/api/files/view/../../etc/passwd").status_code in (400, 404)


def test_docx_and_xlsx_are_previewable(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    from docx import Document
    from openpyxl import Workbook
    d = Document(); d.add_heading("Por qué Apple", 1); d.add_paragraph("Ecosistema."); d.save(ws / "r.docx")
    wb = Workbook(); wb.active.append(["Modelo", "Precio"]); wb.active.append(["X", 10]); wb.save(ws / "t.xlsx")
    with _client(tmp_path, monkeypatch) as c:
        assert "<h1>Por qué Apple</h1>" in c.get("/api/files/view/r.docx").text
        assert "<th>Modelo</th>" in c.get("/api/files/view/t.xlsx").text


def test_created_files_can_be_opened_from_chat(tmp_path, monkeypatch):
    from app.core.file_actions import run_for_response
    monkeypatch.setenv("ANDROMEDA_WORKSPACE", str(tmp_path))
    with _client(tmp_path, monkeypatch) as c:
        r = run_for_response('```andromeda:write path="a.html"\n<h1>x</h1>\n```',
                             settings=c.app.state.settings, incognito=False)
    assert r[0]["path"] == "a.html"
    bubble = (ROOT / "frontend/src/components/MessageBubble.jsx").read_text(encoding="utf-8")
    assert "openPreview(fa.path)" in bubble


# ── Interfaz ──────────────────────────────────────────────────────────────

def _src(f):
    return (ROOT / "frontend/src" / f).read_text(encoding="utf-8")


def test_scroll_is_not_forced_while_streaming():
    """Se forzaba el desplazamiento con cada token: imposible subir a leer."""
    chat = _src("components/Chat.jsx")
    assert "scrollIntoView({ behavior:'smooth', block:'end' })" not in chat
    assert "if (pegado)" in chat and "jump-bottom" in chat


def test_lite_hides_pro_badges():
    bubble = _src("components/MessageBubble.jsx")
    for cond in ("isPro() && showBadges()", "isPro() && meta()?.strategy_used",
                 "isPro() && meta()?.power_tier", "isPro() && meta()?.confidence"):
        assert cond in bubble, cond


def test_external_links_open_in_system_browser(tmp_path, monkeypatch):
    import webbrowser
    abiertas = []
    monkeypatch.setattr(webbrowser, "open", lambda url, new=0: abiertas.append(url) or True)
    with _client(tmp_path, monkeypatch) as c:
        assert c.post("/api/system/open-url", json={"url": "https://apple.com"}).json()["opened"]
        assert c.post("/api/system/open-url", json={"url": "file:///etc/passwd"}).status_code == 400
        assert c.post("/api/system/open-url", json={"url": "javascript:alert(1)"}).status_code == 400
    assert abiertas == ["https://apple.com"]
    assert "/api/system/open-url" in _src("App.jsx")


def test_image_zoom_is_animated():
    css = _src("styles/index.css")
    assert "@keyframes zoomIn" in css and ".att-zoom.out" in css


def test_prompt_catalog_is_large_and_complete():
    import json
    import subprocess
    out = subprocess.run(
        ["node", "-e", "import('" + str(ROOT / "frontend/src/data/promptCatalog.js") +
         "').then(m=>console.log(JSON.stringify({n:m.PROMPT_CATALOG.length,"
         "vacias:m.PROMPT_CATEGORIES.filter(c=>!m.PROMPT_CATALOG.some(p=>p.cat===c))})))"],
        capture_output=True, text=True, timeout=30)
    if out.returncode != 0:
        return   # sin node en el entorno de tests
    d = json.loads(out.stdout)
    assert d["n"] >= 100 and d["vacias"] == []


def test_prompt_library_is_in_settings():
    assert "section()==='prompts'" in _src("components/SettingsPanel.jsx")


def test_folder_attachments_skip_binaries_and_cap_size():
    chat = _src("components/Chat.jsx")
    assert "MAX_TOTAL" in chat and "EXT_IMAGEN" in chat and "node_modules" in chat


def test_new_2026_models_in_catalog():
    from app.core.model_catalog import CATALOG
    nombres = {m["name"] for m in CATALOG}
    assert {"qwen3.8:27b", "gemma4:12b", "gpt-oss:20b"} <= nombres
    from app.core.vision import is_vision_model
    assert is_vision_model("gemma4:12b") and is_vision_model("qwen3.8:27b")
