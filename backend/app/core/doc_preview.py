"""
doc_preview.py — Vista previa en HTML de documentos que el navegador no muestra.

Word, Excel y CSV no se pueden ver en un iframe tal cual: antes, al pulsar sobre
un .docx en Artefactos se intentaba abrir una ventana nueva, que en la app de
escritorio no hace nada. Aquí se convierten a un HTML sencillo y legible.
"""
from __future__ import annotations

import csv
import html
import io
from pathlib import Path

_ESTILO = """<style>
body{font:15px/1.6 -apple-system,BlinkMacSystemFont,'Segoe UI',Inter,sans-serif;color:#1f2330;
     max-width:860px;margin:32px auto;padding:0 28px;background:#fff}
h1,h2,h3{line-height:1.25;margin:1.4em 0 .5em}h1{font-size:1.9em}h2{font-size:1.45em}h3{font-size:1.2em}
table{border-collapse:collapse;margin:14px 0;font-size:14px;width:100%}
td,th{border:1px solid #d8dbe4;padding:6px 10px;text-align:left;vertical-align:top}
th{background:#f3f4f8;font-weight:600}.hoja{font-weight:600;margin-top:28px;color:#555}
li{margin:.2em 0}
</style>"""


def _pagina(cuerpo: str) -> str:
    return f"<!doctype html><html><head><meta charset='utf-8'>{_ESTILO}</head><body>{cuerpo}</body></html>"


def docx_to_html(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    partes, en_lista = [], False
    for bloque in doc.element.body.iterchildren():
        etiqueta = bloque.tag.rsplit("}", 1)[-1]
        if etiqueta == "p":
            from docx.text.paragraph import Paragraph
            p = Paragraph(bloque, doc)
            texto = html.escape(p.text)
            estilo = (p.style.name or "").lower() if p.style is not None else ""
            es_lista = "list" in estilo or "lista" in estilo
            if es_lista and not en_lista:
                partes.append("<ul>"); en_lista = True
            if not es_lista and en_lista:
                partes.append("</ul>"); en_lista = False
            if not texto:
                continue
            if estilo.startswith(("heading 1", "título 1", "title")):
                partes.append(f"<h1>{texto}</h1>")
            elif estilo.startswith(("heading 2", "título 2")):
                partes.append(f"<h2>{texto}</h2>")
            elif estilo.startswith(("heading", "título")):
                partes.append(f"<h3>{texto}</h3>")
            elif es_lista:
                partes.append(f"<li>{texto}</li>")
            else:
                partes.append(f"<p>{texto}</p>")
        elif etiqueta == "tbl":
            from docx.table import Table
            t = Table(bloque, doc)
            filas = "".join("<tr>" + "".join(f"<td>{html.escape(c.text)}</td>" for c in r.cells)
                            + "</tr>" for r in t.rows)
            partes.append(f"<table>{filas}</table>")
    if en_lista:
        partes.append("</ul>")
    return _pagina("\n".join(partes) or "<p><em>(documento vacío)</em></p>")


def xlsx_to_html(path: Path, max_filas: int = 500) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    partes = []
    for hoja in wb.worksheets:
        filas = []
        for i, fila in enumerate(hoja.iter_rows(values_only=True)):
            if i >= max_filas:
                filas.append(f"<tr><td colspan='99'>… (más de {max_filas} filas)</td></tr>")
                break
            celdas = "".join(f"<{'th' if i == 0 else 'td'}>{html.escape('' if v is None else str(v))}"
                             f"</{'th' if i == 0 else 'td'}>" for v in fila)
            filas.append(f"<tr>{celdas}</tr>")
        partes.append(f"<div class='hoja'>{html.escape(hoja.title)}</div><table>{''.join(filas)}</table>")
    return _pagina("".join(partes) or "<p><em>(libro vacío)</em></p>")


def csv_to_html(path: Path, max_filas: int = 500) -> str:
    texto = path.read_text(encoding="utf-8", errors="replace")
    filas = []
    for i, fila in enumerate(csv.reader(io.StringIO(texto))):
        if i >= max_filas:
            break
        tag = "th" if i == 0 else "td"
        filas.append("<tr>" + "".join(f"<{tag}>{html.escape(c)}</{tag}>" for c in fila) + "</tr>")
    return _pagina(f"<table>{''.join(filas)}</table>")


CONVERTIBLES = {".docx": docx_to_html, ".xlsx": xlsx_to_html, ".csv": csv_to_html}
